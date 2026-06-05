import copy
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any, Iterable, Optional

from fastapi import HTTPException
from pptx import Presentation


SECTION_HEADING_RE = re.compile(r"^\s*#{1,6}\s*(?:(\d+)[.)]?\s*)?(.*?)\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
METRIC_RE = re.compile(
    r"(?<![\w.])[+-]?\$?\d[\d,]*(?:\.\d+)?%?"
    r"(?:\s?(?:pp|bps|ms|sec|secs|seconds|m|k|K|M|B))?(?![\w.])"
)
DATE_CAPTION_LABEL_RE = re.compile(
    r"^(?:selected\s+period|prior\s+period|date\s+range|date\s+context|"
    r"half[-\s]?period\s+definitions?|date\s+caption(?:\s+wrappers?)?)\b"
    r"(?::|\s|$)",
    re.IGNORECASE,
)
DATE_CAPTION_TAG_RE = re.compile(
    r"</?(?:date|period)[_\s-]*(?:caption|context|range|wrapper)s?>",
    re.IGNORECASE,
)
DATE_CAPTION_WRAPPER_PREFIX_RE = re.compile(
    r"^date\s+caption(?:\s+wrappers?)?\s*:\s*",
    re.IGNORECASE,
)
DATE_CAPTION_TRAILING_GUIDANCE_RE = re.compile(
    r"\s*[.;]?\s+keep\s+this\s+near\b.*$",
    re.IGNORECASE,
)
PLACEHOLDER_ENTITY_NOUN_RE = (
    r"(?:Retailer|retailer|RETAILER|Campaign|campaign|CAMPAIGN|Source|source|SOURCE)"
)
PLACEHOLDER_TOKEN_RE = (
    r"(?:[ABCXYZ]|Alpha|Beta|Gamma|Delta|ALPHA|BETA|GAMMA|DELTA|"
    r"Placeholder|Sample|Example|Generic|"
    r"PLACEHOLDER|SAMPLE|EXAMPLE|GENERIC)"
)
GENERIC_PLACEHOLDER_LABEL_RE = re.compile(
    rf"\b(?P<label>{PLACEHOLDER_ENTITY_NOUN_RE}\s+{PLACEHOLDER_TOKEN_RE}|"
    rf"(?:Placeholder|Sample|Example|Generic|PLACEHOLDER|SAMPLE|EXAMPLE|GENERIC)"
    rf"\s+{PLACEHOLDER_ENTITY_NOUN_RE})\b"
)


@dataclass
class ContractIssue:
    reason: str
    message: str
    severity: str = "error"
    stage: str = "contract"
    section_index: Optional[int] = None
    expected: Any = None
    actual: Any = None
    details: Optional[dict[str, Any]] = None

    def to_dict(self) -> dict[str, Any]:
        result = {
            "severity": self.severity,
            "reason": self.reason,
            "message": self.message,
            "stage": self.stage,
        }
        if self.section_index is not None:
            result["section_index"] = self.section_index
        if self.expected is not None:
            result["expected"] = self.expected
        if self.actual is not None:
            result["actual"] = self.actual
        if self.details is not None:
            result["details"] = self.details
        return result


@dataclass
class ContractTable:
    headers: list[str]
    rows: list[list[str]]
    slide_index: Optional[int] = None
    section_title: Optional[str] = None
    required: bool = True
    source: str = "markdown"

    @property
    def column_count(self) -> int:
        return len(self.headers)

    @property
    def row_count(self) -> int:
        return len(self.rows)


@dataclass
class ContractTextBlock:
    text: str
    slide_index: Optional[int] = None
    section_title: Optional[str] = None
    source: str = "locked_text"


@dataclass
class ContractSection:
    index: int
    title: str
    markdown: str
    tables: list[ContractTable] = field(default_factory=list)


@dataclass
class GenerationContractState:
    enabled: bool
    violation_policy: str = "fail"
    slides_markdown: list[str] = field(default_factory=list)
    sections: list[ContractSection] = field(default_factory=list)
    locked_text: list[str] = field(default_factory=list)
    forbidden_additions: list[str] = field(default_factory=list)
    exact_terms: list[str] = field(default_factory=list)
    evidence_tables: list[ContractTable] = field(default_factory=list)
    tables_are_evidence: bool = False


def strict_contract_enabled(request: Any) -> bool:
    return (
        getattr(request, "contract_mode", "off") == "strict"
        or getattr(request, "generation_mode", "standard") == "layout_from_contract"
        or getattr(request, "content_generation", "generate") == "preserve"
    )


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split()).strip()


def _norm(value: Any) -> str:
    return _clean_text(value).lower()


def _strings(values: Iterable[Any]) -> list[str]:
    result = []
    for value in values or []:
        text = _clean_text(value)
        if text:
            result.append(text)
    return result


def _generic_placeholder_labels(values: Iterable[Any]) -> list[str]:
    labels = []
    seen = set()
    for value in values or []:
        for segment in str(value or "").splitlines():
            text = _clean_text(segment)
            if not text:
                continue
            for match in GENERIC_PLACEHOLDER_LABEL_RE.finditer(text):
                label = _clean_text(match.group("label"))
                key = _norm(label)
                if key in seen:
                    continue
                seen.add(key)
                labels.append(label)
    return labels


def _generic_placeholder_issues(
    values: Iterable[Any],
    *,
    stage: str,
) -> list[ContractIssue]:
    labels = _generic_placeholder_labels(values)
    if not labels:
        return []
    return [
        ContractIssue(
            reason="generic_placeholder_content",
            message="Strict contract output contains obvious synthetic placeholder labels.",
            stage=stage,
            actual=labels,
        )
    ]


def _contract_obj(contract: Any, name: str, default: Any) -> Any:
    if contract is None:
        return default
    if isinstance(contract, dict):
        return contract.get(name, default)
    return getattr(contract, name, default)


def _parse_table_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    cells = []
    current = []
    escaped = False
    for char in raw:
        if escaped:
            current.append("|" if char == "|" else f"\\{char}")
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    if escaped:
        current.append("\\")
    cells.append("".join(current).strip())
    return [_clean_text(cell) for cell in cells]


def _strip_html_comments(markdown: str) -> str:
    text = str(markdown or "")
    result: list[str] = []
    i = 0
    while i < len(text):
        start = text.find("<!--", i)
        if start == -1:
            _append_visible_markdown_segment(result, text[i:])
            break
        _append_visible_markdown_segment(result, text[i:start])
        end = text.find("-->", start + 4)
        if end == -1:
            break
        i = end + 3
    return "".join(result)


def _append_visible_markdown_segment(result: list[str], segment: str) -> None:
    if not segment:
        return
    if (
        result
        and result[-1]
        and result[-1][-1] in " \t"
        and segment[0] in " \t"
    ):
        segment = segment.lstrip(" \t")
    result.append(segment)


def _sanitized_markdown_lines(markdown: str) -> list[str]:
    return _strip_html_comments(markdown).splitlines()


def parse_markdown_tables(markdown: str) -> list[ContractTable]:
    lines = _sanitized_markdown_lines(markdown)
    tables: list[ContractTable] = []
    i = 0
    while i < len(lines) - 1:
        if "|" not in lines[i] or not TABLE_SEPARATOR_RE.match(lines[i + 1] or ""):
            i += 1
            continue
        headers = _parse_table_row(lines[i])
        rows: list[list[str]] = []
        i += 2
        while i < len(lines) and "|" in lines[i] and lines[i].strip():
            row = _parse_table_row(lines[i])
            if row:
                rows.append(row)
            i += 1
        if headers and rows:
            tables.append(ContractTable(headers=headers, rows=rows))
        continue
    return tables


def _section_title(markdown: str, fallback_index: int) -> str:
    for line in _sanitized_markdown_lines(markdown):
        match = SECTION_HEADING_RE.match(line)
        if match:
            return _clean_text(match.group(2)) or f"Section {fallback_index}"
    return f"Section {fallback_index}"


def _sections_from_markdown(slides_markdown: list[str]) -> list[ContractSection]:
    sections = []
    for index, markdown in enumerate(slides_markdown, start=1):
        tables = parse_markdown_tables(markdown)
        for table in tables:
            table.slide_index = index
            table.section_title = _section_title(markdown, index)
        sections.append(
            ContractSection(
                index=index,
                title=_section_title(markdown, index),
                markdown=str(markdown or ""),
                tables=tables,
            )
        )
    return sections


def _required_sections_to_markdown(required_sections: list[Any]) -> list[str]:
    markdown = []
    for offset, section in enumerate(required_sections or [], start=1):
        if isinstance(section, str):
            text = section
            title = _section_title(text, offset)
        elif isinstance(section, dict):
            title = _clean_text(section.get("title")) or f"Section {offset}"
            text = _clean_text(section.get("content"))
        else:
            title = _clean_text(getattr(section, "title", "")) or f"Section {offset}"
            text = _clean_text(getattr(section, "content", ""))
        if text and re.search(r"^\s*#{1,6}\s*", text, re.MULTILINE):
            markdown.append(text)
        else:
            markdown.append(f"### {offset}. {title}\n\n{text}".strip())
    return markdown


def _table_from_contract(value: Any) -> Optional[ContractTable]:
    if value is None:
        return None
    if isinstance(value, dict):
        headers = value.get("headers") or value.get("columns") or []
        rows = value.get("rows") or []
        markdown = value.get("markdown")
        slide_index = value.get("slide_index") or value.get("slideIndex")
        section_title = value.get("section_title") or value.get("sectionTitle")
        required = value.get("required", True)
    else:
        headers = getattr(value, "headers", []) or []
        rows = getattr(value, "rows", []) or []
        markdown = getattr(value, "markdown", None)
        slide_index = getattr(value, "slide_index", None)
        section_title = getattr(value, "section_title", None)
        required = getattr(value, "required", True)
    if markdown and not headers and not rows:
        parsed = parse_markdown_tables(markdown)
        if parsed:
            table = parsed[0]
            table.slide_index = slide_index
            table.section_title = section_title
            table.required = bool(required)
            table.source = "generation_contract"
            return table
    clean_headers = _strings(headers)
    clean_rows = [[_clean_text(cell) for cell in row] for row in rows or []]
    if not clean_headers or not clean_rows:
        return None
    return ContractTable(
        headers=clean_headers,
        rows=clean_rows,
        slide_index=slide_index,
        section_title=section_title,
        required=bool(required),
        source="generation_contract",
    )


def _table_signature(table: ContractTable) -> tuple[Any, ...]:
    location = (
        table.slide_index
        if table.slide_index is not None
        else _norm(table.section_title)
    )
    return (
        location,
        tuple(_norm(cell) for cell in table.headers),
        tuple(tuple(_norm(cell) for cell in row) for row in table.rows),
    )


def _text_signature(block: ContractTextBlock) -> tuple[Any, ...]:
    location = (
        block.slide_index
        if block.slide_index is not None
        else _norm(block.section_title)
    )
    return location, _norm(block.text)


def _dedupe_tables(tables: Iterable[ContractTable]) -> list[ContractTable]:
    seen: set[tuple[Any, ...]] = set()
    result = []
    for table in tables:
        signature = _table_signature(table)
        if signature in seen:
            continue
        seen.add(signature)
        result.append(table)
    return result


def _dedupe_text_blocks(blocks: Iterable[ContractTextBlock]) -> list[ContractTextBlock]:
    seen: set[tuple[Any, ...]] = set()
    result = []
    for block in blocks:
        signature = _text_signature(block)
        if signature in seen:
            continue
        seen.add(signature)
        result.append(block)
    return result


def _source_text(slides_markdown: list[str]) -> str:
    return "\n".join(_strip_html_comments(markdown) for markdown in slides_markdown or [])


def _extract_source_metric_terms(slides_markdown: list[str]) -> set[str]:
    terms = set()
    for match in METRIC_RE.findall(_source_text(slides_markdown)):
        cleaned = _clean_text(match)
        if cleaned and not cleaned.isdigit():
            terms.add(cleaned)
    return terms


def _normalize_date_caption_line(line: str) -> str:
    text = _clean_text(line)
    text = re.sub(r"^\s*#{1,6}\s*", "", text)
    text = re.sub(r"^\s*(?:[-*+]\s*|\d+[.)]\s*|>\s*)+", "", text)
    text = DATE_CAPTION_TAG_RE.sub("", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"\s+[-\u2013\u2014]\s+", " - ", text)
    text = re.sub(r"\s*:\s*", ": ", text)
    text = DATE_CAPTION_WRAPPER_PREFIX_RE.sub("", text)
    text = DATE_CAPTION_TRAILING_GUIDANCE_RE.sub("", text)
    return _clean_text(text).strip("*_ ")


def _is_date_caption_term(caption: str) -> bool:
    if not (
        DATE_CAPTION_LABEL_RE.match(caption)
        or re.match(r"^(?:selected|prior)\s*:", caption, re.IGNORECASE)
    ):
        return False
    if caption.rstrip().endswith(":"):
        return False
    return ":" in caption or bool(re.search(r"\b20\d{2}\b", caption))


def _extract_source_date_caption_terms(slides_markdown: list[str]) -> set[str]:
    terms = set()
    for line in _source_text(slides_markdown).splitlines():
        caption = _normalize_date_caption_line(line)
        if caption and _is_date_caption_term(caption):
            terms.add(caption)
    return terms


def _extract_source_exact_terms(slides_markdown: list[str]) -> list[str]:
    terms = _extract_source_metric_terms(slides_markdown)
    terms.update(_extract_source_date_caption_terms(slides_markdown))
    return sorted(terms, key=lambda item: (len(item), item.lower()))


def build_generation_contract_state(request: Any) -> GenerationContractState:
    enabled = strict_contract_enabled(request)
    contract = getattr(request, "generation_contract", None)
    violation_policy = _contract_obj(contract, "violation_policy", "fail")
    slides_markdown = list(getattr(request, "slides_markdown", None) or [])
    required_sections = list(_contract_obj(contract, "required_sections", []) or [])
    if enabled and not slides_markdown and required_sections:
        slides_markdown = _required_sections_to_markdown(required_sections)

    sections = _sections_from_markdown(slides_markdown)
    evidence_tables = [
        table
        for table in (_table_from_contract(value) for value in _contract_obj(contract, "evidence_tables", []) or [])
        if table
    ]
    for section in sections:
        evidence_tables.extend(section.tables)
    evidence_tables = _dedupe_tables(evidence_tables)

    exact_terms = _strings(_contract_obj(contract, "exact_terms", []) or [])
    if enabled:
        exact_terms.extend(term for term in _extract_source_exact_terms(slides_markdown) if term not in exact_terms)

    return GenerationContractState(
        enabled=enabled,
        violation_policy=violation_policy,
        slides_markdown=slides_markdown,
        sections=sections,
        locked_text=_strings(_contract_obj(contract, "locked_text", []) or []),
        forbidden_additions=_strings(_contract_obj(contract, "forbidden_additions", []) or []),
        exact_terms=exact_terms,
        evidence_tables=evidence_tables,
        tables_are_evidence=bool(_contract_obj(contract, "tables_are_evidence", False)),
    )


def diagnostic_payload(
    issues: list[ContractIssue],
    *,
    stage: str,
    status: str = "fail",
) -> dict[str, Any]:
    return {
        "reason": "generation_contract_violation",
        "status": status,
        "stage": stage,
        "issues": [issue.to_dict() for issue in issues],
    }


def enforce_contract_or_raise(
    state: GenerationContractState,
    issues: list[ContractIssue],
    *,
    stage: str,
) -> None:
    if not state.enabled or not issues:
        return
    payload = diagnostic_payload(issues, stage=stage)
    if state.violation_policy == "warn":
        print(f"[generation_contract] warning: {payload}", flush=True)
        return
    raise HTTPException(status_code=422, detail=payload)


def validate_contract_request(state: GenerationContractState) -> list[ContractIssue]:
    if not state.enabled:
        return []
    issues: list[ContractIssue] = []
    if not state.sections:
        issues.append(
            ContractIssue(
                reason="missing_required_sections",
                message="Strict contract mode requires slides_markdown or generation_contract.required_sections.",
                stage="request",
            )
        )
        return issues
    source_text = "\n\n".join(
        _strip_html_comments(section.markdown) for section in state.sections
    )
    for locked in state.locked_text:
        if locked not in source_text:
            issues.append(
                ContractIssue(
                    reason="missing_locked_text",
                    message="Locked text is not present in the supplied source contract.",
                    stage="request",
                    expected=locked,
                )
            )
    for table in state.evidence_tables:
        if table.column_count == 0 or table.row_count == 0:
            issues.append(
                ContractIssue(
                    reason="empty_evidence_table",
                    message="Evidence table must include headers and at least one row.",
                    stage="request",
                    section_index=table.slide_index,
                )
            )
    return issues


def validate_structure(
    state: GenerationContractState,
    slide_count: int,
    *,
    stage: str,
) -> list[ContractIssue]:
    if not state.enabled:
        return []
    expected = len(state.sections)
    if slide_count != expected:
        if slide_count < expected:
            reason = "skipped_section"
            message = "Strict mode must preserve every numbered section as a slide."
        else:
            reason = "extra_slide"
            message = "Strict mode must not add slides beyond the numbered sections."
        return [
            ContractIssue(
                reason=reason,
                message=message,
                stage=stage,
                expected=expected,
                actual=slide_count,
            )
        ]
    return []


def contract_instructions_for_slide(
    state: GenerationContractState,
    slide_index: int,
) -> str:
    if not state.enabled:
        return ""
    section = state.sections[slide_index] if slide_index < len(state.sections) else None
    lines = [
        "# Strict Contract Preservation",
        "You may choose hierarchy, spacing, and concise non-critical wording, but you must not change facts.",
        "Preserve exact metrics, dates, labels, and evidence nouns from the slide content.",
        "Do not add advice, follow-up actions, or unsupported claims that are not in the slide content.",
    ]
    if section:
        lines.append(f"This is section {section.index}: {section.title}. Keep it as this slide only.")
    visible_section_text = _strip_html_comments(section.markdown) if section else ""
    locked_for_slide = [
        locked
        for locked in state.locked_text
        if section and locked in visible_section_text
    ]
    if locked_for_slide:
        lines.append("Locked text to copy verbatim:")
        lines.extend(f"- {locked}" for locked in locked_for_slide)
    table_count = len(section.tables) if section else 0
    if table_count:
        lines.append("Render markdown tables as table/data fields, not prose.")
    if state.forbidden_additions:
        lines.append("Forbidden additions:")
        lines.extend(f"- {item}" for item in state.forbidden_additions)
    return "\n".join(lines)


TEXT_FIELD_KEY_PRIORITY = (
    ("body", "description", "paragraph", "content", "summary", "narrative", "details", "insight"),
    ("subtitle", "caption", "note", "copy"),
    ("title", "heading", "headline"),
)
TEXT_FIELD_EXCLUDED_KEYS = {
    "__speaker_note__",
    "__image_prompt__",
    "__image_url__",
    "__icon_query__",
    "__icon_url__",
    "headers",
    "columns",
    "rows",
    "label",
    "labels",
    "value",
    "values",
    "metric",
    "metrics",
}
TEXT_FIELD_EXCLUDED_PATH_PARTS = {
    "table",
    "tabledata",
    "chart",
    "chartdata",
    "image",
    "icon",
    "logo",
    "palette",
    "series",
    "axis",
    "data",
}
VISIBLE_JSON_TEXT_EXCLUDED_KEYS = {
    "__speaker_note__",
    "__image_prompt__",
    "__image_url__",
    "__icon_query__",
    "__icon_url__",
}
VISIBLE_JSON_TEXT_EXCLUDED_PATH_PARTS = {
    "image",
    "icon",
    "logo",
}
RENDERER_INSTRUCTION_RE = re.compile(
    r"^\s*(?:speaker\s+notes?|presenter\s+notes?|render(?:er)?\s+instructions?|"
    r"layout\s+instructions?|design\s+instructions?|image\s+prompt|icon\s+query)\s*:",
    re.IGNORECASE,
)
PRESERVED_TITLE_LINE_RE = re.compile(r"^\s*title\s*:\s*(.*?)\s*$", re.IGNORECASE)
PRESERVED_SUBTITLE_LINE_RE = re.compile(
    r"^\s*(subtitle|period)\s*:\s*(.*?)\s*$",
    re.IGNORECASE,
)


def _key_priority(key: str) -> Optional[int]:
    normalized = re.sub(r"[^a-z0-9]+", "", key.lower())
    if normalized in TEXT_FIELD_EXCLUDED_KEYS:
        return None
    for priority, names in enumerate(TEXT_FIELD_KEY_PRIORITY):
        if any(name in normalized for name in names):
            return priority
    return 5


def _path_allows_visible_text(path: list[str]) -> bool:
    for part in path:
        normalized = re.sub(r"[^a-z0-9_]+", "", part.lower())
        if normalized in TEXT_FIELD_EXCLUDED_KEYS:
            return False
        if normalized.startswith("__"):
            return False
        if any(excluded in normalized for excluded in TEXT_FIELD_EXCLUDED_PATH_PARTS):
            return False
    return True


def _path_allows_visible_json_text(path: list[str]) -> bool:
    for part in path:
        normalized = re.sub(r"[^a-z0-9_]+", "", part.lower())
        if normalized in VISIBLE_JSON_TEXT_EXCLUDED_KEYS:
            return False
        if normalized.startswith("__"):
            return False
        if any(excluded in normalized for excluded in VISIBLE_JSON_TEXT_EXCLUDED_PATH_PARTS):
            return False
    return True


def _is_string_schema(schema: Any) -> bool:
    if not isinstance(schema, dict):
        return False
    schema_type = schema.get("type")
    return schema_type == "string" or "maxLength" in schema or "minLength" in schema


def _max_length(schema: dict) -> Optional[int]:
    value = schema.get("maxLength") if isinstance(schema, dict) else None
    return value if isinstance(value, int) else None


def _schema_text_paths(
    schema: Any,
    path: Optional[list[str]] = None,
) -> list[tuple[list[str], int]]:
    path = path or []
    if not isinstance(schema, dict):
        return []

    props = schema.get("properties")
    paths: list[tuple[list[str], int]] = []
    if isinstance(props, dict):
        for key, child in props.items():
            child_path = [*path, key]
            if not _path_allows_visible_text(child_path):
                continue
            priority = _key_priority(key)
            if priority is not None and _is_string_schema(child):
                paths.append((child_path, priority))
            paths.extend(_schema_text_paths(child, child_path))
    return sorted(paths, key=lambda item: (item[1], len(item[0]), ".".join(item[0])))


def _schema_table_paths(schema: Any, path: Optional[list[str]] = None) -> list[tuple[list[str], str]]:
    path = path or []
    if not isinstance(schema, dict):
        return []
    props = schema.get("properties")
    paths: list[tuple[list[str], str]] = []
    if isinstance(props, dict):
        row_schema = props.get("rows")
        if "headers" in props and row_schema is not None:
            paths.append((path, "headers"))
        if "columns" in props and row_schema is not None:
            paths.append((path, "columns"))
        for key, child in props.items():
            paths.extend(_schema_table_paths(child, [*path, key]))
    return paths


def _schema_at_path(schema: dict, path: list[str]) -> dict:
    current = schema
    for key in path:
        current = current.get("properties", {}).get(key, {})
    return current if isinstance(current, dict) else {}


def _max_items(schema: dict, key: str) -> Optional[int]:
    child = schema.get("properties", {}).get(key, {})
    value = child.get("maxItems") if isinstance(child, dict) else None
    return value if isinstance(value, int) else None


def _min_items(schema: dict, key: str) -> Optional[int]:
    child = schema.get("properties", {}).get(key, {})
    value = child.get("minItems") if isinstance(child, dict) else None
    return value if isinstance(value, int) else None


def _schema_int(schema: Any, key: str) -> Optional[int]:
    value = schema.get(key) if isinstance(schema, dict) else None
    return value if isinstance(value, int) else None


def _row_items_schema(schema: dict) -> dict:
    rows_schema = schema.get("properties", {}).get("rows", {})
    items_schema = rows_schema.get("items") if isinstance(rows_schema, dict) else None
    return items_schema if isinstance(items_schema, dict) else {}


def _relax_min_items(schema: Any, item_count: int) -> None:
    if not isinstance(schema, dict):
        return
    value = schema.get("minItems")
    if isinstance(value, int) and value > item_count:
        schema["minItems"] = item_count


def _relax_table_schema_min_items(
    schema: dict,
    header_key: str,
    table: ContractTable,
) -> None:
    props = schema.get("properties", {})
    if not isinstance(props, dict):
        return

    header_schema = props.get(header_key)
    _relax_min_items(header_schema, table.column_count)

    rows_schema = props.get("rows")
    _relax_min_items(rows_schema, table.row_count)
    if isinstance(rows_schema, dict):
        row_item_schema = rows_schema.get("items")
        _relax_min_items(row_item_schema, table.column_count)


def _get_path(target: dict, path: list[str]) -> Any:
    current: Any = target
    for key in path:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def _set_path(target: dict, path: list[str], value: Any) -> None:
    current = target
    for key in path[:-1]:
        child = current.get(key)
        if not isinstance(child, dict):
            child = {}
            current[key] = child
        current = child
    if path:
        current[path[-1]] = value
    else:
        target.clear()
        target.update(value)


def _max_table_cells_per_row(table: ContractTable) -> int:
    return max([table.column_count, *(len(row) for row in table.rows)], default=0)


def _min_table_cells_per_row(table: ContractTable) -> int:
    return min([table.column_count, *(len(row) for row in table.rows)], default=0)


def _table_fit_candidate(
    table: ContractTable,
    schema: dict,
    header_key: str,
    path: Optional[list[str]] = None,
) -> dict[str, Any]:
    candidate: dict[str, Any] = {
        "header_key": header_key,
    }
    if path is not None:
        candidate["field_path"] = _field_path(path)

    max_cols = _max_items(schema, header_key)
    min_cols = _min_items(schema, header_key)
    max_rows = _max_items(schema, "rows")
    min_rows = _min_items(schema, "rows")
    row_items_schema = _row_items_schema(schema)
    max_cells_per_row = _schema_int(row_items_schema, "maxItems")
    min_cells_per_row = _schema_int(row_items_schema, "minItems")
    bounds = {
        "maxColumns": max_cols,
        "minColumns": min_cols,
        "maxRows": max_rows,
        "minRows": min_rows,
        "maxCellsPerRow": max_cells_per_row,
        "minCellsPerRow": min_cells_per_row,
    }
    candidate.update({key: value for key, value in bounds.items() if value is not None})

    failures: list[str] = []
    if max_cols is not None and table.column_count > max_cols:
        failures.append("columns_exceed_maxColumns")
    if max_rows is not None and table.row_count > max_rows:
        failures.append("rows_exceed_maxRows")
    if (
        max_cells_per_row is not None
        and _max_table_cells_per_row(table) > max_cells_per_row
    ):
        failures.append("cells_exceed_maxCellsPerRow")
    if min_cells_per_row is not None:
        relaxed_min_cells_per_row = min(min_cells_per_row, table.column_count)
        if _min_table_cells_per_row(table) < relaxed_min_cells_per_row:
            failures.append("cells_below_minCellsPerRow")

    candidate["fits"] = not failures
    if failures:
        candidate["failure_reasons"] = failures
    return candidate


def _table_fits_schema(table: ContractTable, schema: dict, header_key: str) -> bool:
    return bool(_table_fit_candidate(table, schema, header_key)["fits"])


def _table_source_path(table: ContractTable, slide_index: int) -> str:
    if table.source == "generation_contract":
        return "generation_contract.evidence_tables"
    source_index = table.slide_index - 1 if table.slide_index is not None else slide_index
    return f"slides_markdown[{source_index}]"


def _best_table_fit_candidate(
    candidates: list[dict[str, Any]],
) -> Optional[dict[str, Any]]:
    if not candidates:
        return None
    return min(
        candidates,
        key=lambda candidate: (
            len(candidate.get("failure_reasons", [])),
            candidate.get("field_path", ""),
        ),
    )


def _table_fit_issue_details(
    table: ContractTable,
    slide_schema: dict,
    schema_paths: list[tuple[list[str], str]],
    slide_index: int,
    used_paths: Optional[set[tuple[str, ...]]] = None,
) -> dict[str, Any]:
    used_paths = used_paths or set()
    candidates = [
        _table_fit_candidate(
            table,
            _schema_at_path(slide_schema, path),
            header_key,
            path,
        )
        for path, header_key in schema_paths
        if tuple(path) not in used_paths
    ]
    best = _best_table_fit_candidate(candidates)
    details: dict[str, Any] = {
        "issue_code": "STRICT_PRESERVE_TABLE_FIT_FAILED",
        "content_kind": "evidence_table",
        "source_path": _table_source_path(table, slide_index),
        "slide_index": slide_index,
        "section_index": slide_index + 1,
        "columns": table.column_count,
        "rows": table.row_count,
        "maxCellsInTableRow": _max_table_cells_per_row(table),
        "minCellsInTableRow": _min_table_cells_per_row(table),
        "candidate_count": len(candidates),
    }
    if best is not None:
        details["candidate_path"] = best.get("field_path")
        for key in (
            "maxColumns",
            "minColumns",
            "maxRows",
            "minRows",
            "maxCellsPerRow",
            "minCellsPerRow",
        ):
            if key in best:
                details[key] = best[key]
    if candidates:
        details["candidates"] = candidates[:5]
    return details


def _table_payload(table: ContractTable, header_key: str) -> dict[str, Any]:
    return {
        header_key: list(table.headers),
        "rows": [list(row) for row in table.rows],
    }


def _tables_for_slide(
    state: GenerationContractState,
    slide_index: int,
) -> list[ContractTable]:
    section = state.sections[slide_index] if slide_index < len(state.sections) else None
    return _dedupe_tables(
        [
            *(section.tables if section else []),
            *(
                table
                for table in state.evidence_tables
                if table.slide_index == slide_index + 1
            ),
        ]
    )


def _text_blocks_for_slide(
    state: GenerationContractState,
    slide_index: int,
) -> list[ContractTextBlock]:
    section = state.sections[slide_index] if slide_index < len(state.sections) else None
    if not section:
        return []
    blocks = []
    section_text = _strip_html_comments(section.markdown)
    for locked in state.locked_text:
        if locked in section_text:
            blocks.append(
                ContractTextBlock(
                    text=locked,
                    slide_index=slide_index + 1,
                    section_title=section.title,
                )
            )
    return _dedupe_text_blocks(blocks)


def _combined_locked_text_for_slide(
    state: GenerationContractState,
    slide_index: int,
) -> str:
    return "\n".join(block.text for block in _text_blocks_for_slide(state, slide_index))


def _locked_texts_for_slide(
    state: GenerationContractState,
    slide_index: int,
) -> list[str]:
    return [block.text for block in _text_blocks_for_slide(state, slide_index)]


def _locked_texts_missing_from_visible_content(
    locked_texts: list[str],
    visible_text: str,
) -> list[str]:
    return [text for text in locked_texts if text not in visible_text]


def _text_fits_schema(text: str, schema: dict) -> bool:
    max_length = _max_length(schema)
    return max_length is None or len(text) <= max_length


def _markdown_heading_title_line(markdown: str) -> Optional[tuple[int, str]]:
    for index, line in enumerate(_sanitized_markdown_lines(markdown)):
        match = SECTION_HEADING_RE.match(line)
        if match:
            title = _clean_text(match.group(2))
            if title:
                return index, title
    return None


def _markdown_heading_title(markdown: str) -> Optional[str]:
    title_line = _markdown_heading_title_line(markdown)
    return title_line[1] if title_line else None


def _is_title_path(path: list[str]) -> bool:
    if not path:
        return False
    normalized = re.sub(r"[^a-z0-9]+", "", path[-1].lower())
    if normalized in {"title", "heading", "headline"}:
        return True
    if normalized == "subtitle":
        return False
    return normalized.endswith(("title", "heading", "headline")) or normalized.startswith(
        ("title", "heading", "headline")
    )


def _is_subtitle_path(path: list[str]) -> bool:
    if not path:
        return False
    normalized = re.sub(r"[^a-z0-9]+", "", path[-1].lower())
    return any(name in normalized for name in ("subtitle", "caption", "note"))


def _field_path(path: list[str]) -> str:
    return ".".join(path)


def _filtered_text_paths(
    slide_schema: dict,
    *,
    title: Optional[bool] = None,
    subtitle: Optional[bool] = None,
    excluded_paths: Optional[set[tuple[str, ...]]] = None,
) -> list[list[str]]:
    excluded_paths = excluded_paths or set()
    paths = []
    for path, _ in _schema_text_paths(slide_schema):
        if tuple(path) in excluded_paths:
            continue
        if title is not None and _is_title_path(path) != title:
            continue
        if subtitle is not None and _is_subtitle_path(path) != subtitle:
            continue
        paths.append(path)
    return paths


def _text_fit_candidates(
    slide_schema: dict,
    text: str,
    *,
    title: Optional[bool] = None,
    subtitle: Optional[bool] = None,
    excluded_paths: Optional[set[tuple[str, ...]]] = None,
) -> list[dict[str, Any]]:
    candidates = []
    for path in _filtered_text_paths(
        slide_schema,
        title=title,
        subtitle=subtitle,
        excluded_paths=excluded_paths,
    ):
        schema = _schema_at_path(slide_schema, path)
        candidate: dict[str, Any] = {
            "field_path": _field_path(path),
        }
        max_length = _max_length(schema)
        if max_length is not None:
            candidate["maxLength"] = max_length
            candidate["overBy"] = max(0, len(text) - max_length)
        candidates.append(candidate)
    return candidates


def _best_fit_candidate(candidates: list[dict[str, Any]]) -> Optional[dict[str, Any]]:
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda candidate: (
            candidate.get("maxLength") is None,
            candidate.get("maxLength", 0),
        ),
    )


def _fit_issue_details(
    slide_schema: dict,
    text: str,
    *,
    content_kind: str,
    slide_index: int,
    title: Optional[bool] = None,
    subtitle: Optional[bool] = None,
    excluded_paths: Optional[set[tuple[str, ...]]] = None,
    fallback_title: Optional[bool] = None,
) -> dict[str, Any]:
    candidates = _text_fit_candidates(
        slide_schema,
        text,
        title=title,
        subtitle=subtitle,
        excluded_paths=excluded_paths,
    )
    best = _best_fit_candidate(candidates)
    fallback_candidates: list[dict[str, Any]] = []
    if best is None and fallback_title is not None:
        fallback_candidates = _text_fit_candidates(
            slide_schema,
            text,
            title=fallback_title,
            subtitle=subtitle,
            excluded_paths=excluded_paths,
        )
        best = _best_fit_candidate(fallback_candidates)

    details: dict[str, Any] = {
        "issue_code": "STRICT_PRESERVE_FIT_FAILED",
        "content_kind": content_kind,
        "source_path": f"slides_markdown[{slide_index}]",
        "length": len(text),
        "slide_index": slide_index,
        "section_index": slide_index + 1,
        "candidate_count": len(candidates),
    }
    if best is not None:
        details["field_path"] = best["field_path"]
        if "maxLength" in best:
            details["maxLength"] = best["maxLength"]
    if candidates:
        details["candidates"] = candidates[:5]
    if fallback_candidates:
        details["fallback_candidate_count"] = len(fallback_candidates)
        details["fallback_candidates"] = fallback_candidates[:5]
    return details


def _markdown_table_block_end(lines: list[str], start: int) -> Optional[int]:
    if (
        start + 1 >= len(lines)
        or "|" not in lines[start]
        or not TABLE_SEPARATOR_RE.match(lines[start + 1] or "")
    ):
        return None

    i = start + 2
    has_rows = False
    while i < len(lines) and "|" in lines[i] and lines[i].strip():
        has_rows = True
        i += 1
    if not has_rows:
        return None
    return i


def _preserved_label_line(
    markdown: str,
    pattern: re.Pattern,
) -> Optional[tuple[int, str]]:
    for index, line in enumerate(_sanitized_markdown_lines(markdown)):
        match = pattern.match(line)
        if not match:
            continue
        value = _clean_text(match.group(match.lastindex or 1))
        if value:
            return index, value
    return None


def _preserved_markdown_body_text(
    markdown: str,
    consumed_line_indices: Optional[set[int]] = None,
) -> str:
    lines = _sanitized_markdown_lines(markdown)
    consumed_line_indices = consumed_line_indices or set()
    preserved: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if i in consumed_line_indices:
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("<!--"):
            i += 1
            if "-->" in stripped:
                continue
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            if i < len(lines):
                i += 1
            continue
        table_block_end = _markdown_table_block_end(lines, i)
        if table_block_end is not None:
            i = table_block_end
            continue
        if RENDERER_INSTRUCTION_RE.match(line):
            i += 1
            continue
        heading_match = SECTION_HEADING_RE.match(line)
        if heading_match:
            heading_text = _clean_text(heading_match.group(2))
            if heading_text:
                preserved.append(heading_text)
            i += 1
            continue
        preserved.append(line.rstrip())
        i += 1
    return "\n".join(preserved).strip()


def _first_fitting_text_path(
    slide_schema: dict,
    text: str,
    *,
    title: Optional[bool] = None,
    subtitle: Optional[bool] = None,
    excluded_paths: Optional[set[tuple[str, ...]]] = None,
) -> Optional[list[str]]:
    for path in _filtered_text_paths(
        slide_schema,
        title=title,
        subtitle=subtitle,
        excluded_paths=excluded_paths,
    ):
        if _text_fits_schema(text, _schema_at_path(slide_schema, path)):
            return path
    return None


def _build_preserved_markdown_text_content(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> tuple[dict, list[ContractIssue]]:
    section = state.sections[slide_index] if slide_index < len(state.sections) else None
    markdown = section.markdown if section else ""
    content: dict[str, Any] = {}
    issues: list[ContractIssue] = []
    used_paths: set[tuple[str, ...]] = set()

    consumed_line_indices: set[int] = set()
    pending_title: Optional[str] = None
    title_line = _preserved_label_line(markdown, PRESERVED_TITLE_LINE_RE)
    heading_title_line = _markdown_heading_title_line(markdown)
    title = title_line[1] if title_line else (
        heading_title_line[1] if heading_title_line else None
    )
    if heading_title_line:
        consumed_line_indices.add(heading_title_line[0])
    if title:
        title_path = _first_fitting_text_path(slide_schema, title, title=True)
        if title_path:
            _set_path(content, title_path, title)
            used_paths.add(tuple(title_path))
            if title_line:
                consumed_line_indices.add(title_line[0])
        else:
            pending_title = title
            if title_line:
                consumed_line_indices.add(title_line[0])

    subtitle_line = _preserved_label_line(markdown, PRESERVED_SUBTITLE_LINE_RE)
    if subtitle_line:
        subtitle_path = _first_fitting_text_path(
            slide_schema,
            subtitle_line[1],
            subtitle=True,
            excluded_paths=used_paths,
        )
        if subtitle_path:
            _set_path(content, subtitle_path, subtitle_line[1])
            used_paths.add(tuple(subtitle_path))
            consumed_line_indices.add(subtitle_line[0])

    body_text = _preserved_markdown_body_text(markdown, consumed_line_indices)
    source_text = "\n".join(
        part for part in [pending_title, body_text] if part
    ).strip()
    if source_text:
        body_path = _first_fitting_text_path(
            slide_schema,
            source_text,
            title=False,
            excluded_paths=used_paths,
        )
        if body_path:
            _set_path(content, body_path, source_text)
            used_paths.add(tuple(body_path))
        elif pending_title:
            issues.append(
                ContractIssue(
                    reason="no_compatible_title_layout",
                    message="Selected layout cannot preserve the markdown title in a visible text field.",
                    stage="slide_content",
                    section_index=slide_index + 1,
                    expected=pending_title,
                    details=_fit_issue_details(
                        slide_schema,
                        pending_title,
                        content_kind="preserved_markdown_title",
                        slide_index=slide_index,
                        title=True,
                        excluded_paths=used_paths,
                        fallback_title=False,
                    ),
                )
            )
            if body_text:
                body_path = _first_fitting_text_path(
                    slide_schema,
                    body_text,
                    title=False,
                    excluded_paths=used_paths,
                )
                if body_path:
                    _set_path(content, body_path, body_text)
                    used_paths.add(tuple(body_path))
                else:
                    issues.append(
                        ContractIssue(
                            reason="no_compatible_preserved_text_layout",
                            message="Selected layout cannot preserve the supplied markdown text in a visible text field.",
                            stage="slide_content",
                            section_index=slide_index + 1,
                            expected=body_text,
                            details=_fit_issue_details(
                                slide_schema,
                                body_text,
                                content_kind="preserved_markdown_body",
                                slide_index=slide_index,
                                title=False,
                                excluded_paths=used_paths,
                            ),
                        )
                    )
        else:
            issues.append(
                ContractIssue(
                    reason="no_compatible_preserved_text_layout",
                    message="Selected layout cannot preserve the supplied markdown text in a visible text field.",
                    stage="slide_content",
                    section_index=slide_index + 1,
                    expected=source_text,
                    details=_fit_issue_details(
                        slide_schema,
                        source_text,
                        content_kind="preserved_markdown_body",
                        slide_index=slide_index,
                        title=False,
                        excluded_paths=used_paths,
                    ),
                )
            )

    return content, issues


def build_preserved_slide_content(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> tuple[dict, list[ContractIssue]]:
    content, issues = _build_preserved_markdown_text_content(
        slide_schema,
        state,
        slide_index,
    )
    content, table_issues = overlay_contract_tables(
        content,
        slide_schema,
        state,
        slide_index,
    )
    content, text_issues = overlay_contract_text(
        content,
        slide_schema,
        state,
        slide_index,
    )
    issues.extend([*table_issues, *text_issues])
    return content, issues


def schema_with_contract_table_overrides(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> dict:
    schema = copy.deepcopy(slide_schema or {})
    if not state.enabled:
        return schema

    tables = _tables_for_slide(state, slide_index)
    if not tables:
        return schema

    schema_paths = _schema_table_paths(schema)
    used_paths: set[tuple[str, ...]] = set()
    for table in tables:
        for path, header_key in schema_paths:
            path_key = tuple(path)
            if path_key in used_paths:
                continue
            table_schema = _schema_at_path(schema, path)
            if not _table_fits_schema(table, table_schema, header_key):
                continue
            _relax_table_schema_min_items(table_schema, header_key, table)
            used_paths.add(path_key)
            break
    return schema


def contract_table_issues_for_schema(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> list[ContractIssue]:
    _, issues = overlay_contract_tables({}, slide_schema, state, slide_index)
    return issues


def contract_text_issues_for_schema(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> list[ContractIssue]:
    locked_texts = _locked_texts_for_slide(state, slide_index)
    if not state.enabled or not locked_texts:
        return []
    text = "\n".join(locked_texts)
    for path, _ in _schema_text_paths(slide_schema):
        if _text_fits_schema(text, _schema_at_path(slide_schema, path)):
            return []
    return [
        ContractIssue(
            reason="no_compatible_locked_text_layout",
            message="Selected layout cannot preserve the required locked text in a visible text field.",
            stage="slide_content",
            section_index=slide_index + 1,
            expected=text,
        )
    ]


def contract_layout_issues_for_schema(
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
    *,
    preserve_markdown: bool = False,
) -> list[ContractIssue]:
    if not state.enabled:
        return []
    if not preserve_markdown:
        return [
            *contract_table_issues_for_schema(slide_schema, state, slide_index),
            *contract_text_issues_for_schema(slide_schema, state, slide_index),
        ]

    content, preserved_text_issues = _build_preserved_markdown_text_content(
        slide_schema,
        state,
        slide_index,
    )
    content, table_issues = overlay_contract_tables(
        content,
        slide_schema,
        state,
        slide_index,
    )
    _, text_issues = overlay_contract_text(
        content,
        slide_schema,
        state,
        slide_index,
    )
    return [*preserved_text_issues, *table_issues, *text_issues]


def overlay_contract_tables(
    slide_content: dict,
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> tuple[dict, list[ContractIssue]]:
    if not state.enabled:
        return slide_content, []
    tables = _tables_for_slide(state, slide_index)
    if not tables:
        return slide_content, []

    content = copy.deepcopy(slide_content)
    schema_paths = _schema_table_paths(slide_schema)
    issues: list[ContractIssue] = []
    used_paths: set[tuple[str, ...]] = set()
    for table in tables:
        selected: Optional[tuple[list[str], str]] = None
        for path, header_key in schema_paths:
            path_key = tuple(path)
            if path_key in used_paths:
                continue
            if _table_fits_schema(table, _schema_at_path(slide_schema, path), header_key):
                selected = (path, header_key)
                break
        if not selected:
            issues.append(
                ContractIssue(
                    reason="no_compatible_table_layout",
                    message="Selected layout cannot preserve the required evidence table dimensions.",
                    stage="slide_content",
                    section_index=slide_index + 1,
                    expected={"columns": table.column_count, "rows": table.row_count},
                    details=_table_fit_issue_details(
                        table,
                        slide_schema,
                        schema_paths,
                        slide_index,
                        used_paths,
                    ),
                )
            )
            continue
        path, header_key = selected
        existing = _get_path(content, path)
        payload = _table_payload(table, header_key)
        if path and not isinstance(existing, dict):
            _set_path(content, path, payload)
        elif path:
            _set_path(content, path, {**(existing or {}), **payload})
        else:
            content.update(payload)
        used_paths.add(tuple(path))
    return content, issues


def overlay_contract_text(
    slide_content: dict,
    slide_schema: dict,
    state: GenerationContractState,
    slide_index: int,
) -> tuple[dict, list[ContractIssue]]:
    if not state.enabled:
        return slide_content, []
    locked_texts = _locked_texts_for_slide(state, slide_index)
    if not locked_texts:
        return slide_content, []
    visible_text = visible_text_from_json(slide_content)
    missing_texts = _locked_texts_missing_from_visible_content(
        locked_texts,
        visible_text,
    )
    if not missing_texts:
        return slide_content, []

    text = "\n".join(missing_texts)
    for path, _ in _schema_text_paths(slide_schema):
        if not _text_fits_schema(text, _schema_at_path(slide_schema, path)):
            continue
        existing = _get_path(slide_content, path)
        if existing not in (None, ""):
            continue
        content = copy.deepcopy(slide_content)
        _set_path(content, path, text)
        return content, []

    can_replace_existing_text = not any(
        locked in visible_text for locked in locked_texts
    )
    if can_replace_existing_text:
        for path, _ in _schema_text_paths(slide_schema):
            if not _text_fits_schema(text, _schema_at_path(slide_schema, path)):
                continue
            existing = _get_path(slide_content, path)
            if existing is not None and not isinstance(existing, str):
                continue
            content = copy.deepcopy(slide_content)
            _set_path(content, path, text)
            return content, []

    return slide_content, [
        ContractIssue(
            reason="no_compatible_locked_text_layout",
            message="Selected layout cannot preserve the required locked text in a visible text field.",
            stage="slide_content",
            section_index=slide_index + 1,
            expected=text,
        )
    ]


def _walk_visible_strings(value: Any, path: Optional[list[str]] = None) -> Iterable[str]:
    path = path or []
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for key, child in value.items():
            child_path = [*path, str(key)]
            if not _path_allows_visible_json_text(child_path):
                continue
            yield from _walk_visible_strings(child, child_path)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_visible_strings(child, path)


def visible_text_from_json(value: Any) -> str:
    return "\n".join(_walk_visible_strings(value))


def _table_values_from_json(value: Any) -> list[ContractTable]:
    tables: list[ContractTable] = []
    if isinstance(value, dict):
        headers = value.get("headers") or value.get("columns")
        rows = value.get("rows")
        if isinstance(headers, list) and isinstance(rows, list):
            clean_headers = _strings(headers)
            clean_rows = [[_clean_text(cell) for cell in row] for row in rows if isinstance(row, list)]
            if clean_headers and clean_rows:
                tables.append(ContractTable(headers=clean_headers, rows=clean_rows, source="json"))
        for child in value.values():
            tables.extend(_table_values_from_json(child))
    elif isinstance(value, list):
        for child in value:
            tables.extend(_table_values_from_json(child))
    return tables


def _table_matches(expected: ContractTable, actual: ContractTable) -> bool:
    return (
        [_norm(cell) for cell in expected.headers] == [_norm(cell) for cell in actual.headers]
        and [[_norm(cell) for cell in row] for row in expected.rows]
        == [[_norm(cell) for cell in row] for row in actual.rows]
    )


def validate_slide_json_contract(
    state: GenerationContractState,
    slide_contents: list[dict],
) -> list[ContractIssue]:
    if not state.enabled:
        return []
    issues = validate_structure(state, len(slide_contents), stage="slide_content")
    visible_strings = [
        text for slide in slide_contents for text in _walk_visible_strings(slide)
    ]
    issues.extend(_generic_placeholder_issues(visible_strings, stage="slide_content"))
    full_text = "\n".join(visible_strings)
    normalized_full = _norm(full_text)

    for locked in state.locked_text:
        if locked not in full_text:
            issues.append(
                ContractIssue(
                    reason="missing_locked_text",
                    message="Generated slide JSON did not preserve locked text verbatim.",
                    stage="slide_content",
                    expected=locked,
                )
            )
    for term in state.exact_terms:
        if term and term not in full_text:
            issues.append(
                ContractIssue(
                    reason="changed_metric_date_or_label",
                    message="Generated slide JSON is missing an exact metric, date, or label.",
                    stage="slide_content",
                    expected=term,
                )
            )
    for forbidden in state.forbidden_additions:
        if _norm(forbidden) and _norm(forbidden) in normalized_full:
            issues.append(
                ContractIssue(
                    reason="forbidden_addition",
                    message="Generated slide JSON added forbidden content.",
                    stage="slide_content",
                    expected=forbidden,
                )
            )

    actual_tables = []
    for slide_index, content in enumerate(slide_contents, start=1):
        for table in _table_values_from_json(content):
            table.slide_index = slide_index
            actual_tables.append(table)

    for expected in state.evidence_tables:
        matches = [table for table in actual_tables if _table_matches(expected, table)]
        if not matches and (state.tables_are_evidence or expected.required):
            expected_values = "\n".join(expected.headers + [cell for row in expected.rows for cell in row])
            prose_has_values = all(_norm(cell) in normalized_full for cell in expected_values.splitlines() if _norm(cell))
            issues.append(
                ContractIssue(
                    reason="table_rendered_as_prose" if prose_has_values else "changed_table_values",
                    message="Evidence table was not preserved as a matching table object.",
                    stage="slide_content",
                    section_index=expected.slide_index,
                    expected={"headers": expected.headers, "rows": expected.rows},
                )
            )
    return issues


def _extract_pptx(path: str) -> tuple[str, list[ContractTable], dict[int, str]]:
    prs = Presentation(path)
    texts: list[str] = []
    slide_texts: dict[int, str] = {}
    tables: list[ContractTable] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        current_slide_texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape_text = shape.text
                texts.append(shape_text)
                current_slide_texts.append(shape_text)
            if getattr(shape, "has_table", False):
                rows = []
                headers: list[str] = []
                for row_index, row in enumerate(shape.table.rows):
                    values = [_clean_text(cell.text) for cell in row.cells]
                    texts.extend(values)
                    current_slide_texts.extend(values)
                    if row_index == 0:
                        headers = values
                    else:
                        rows.append(values)
                if headers and rows:
                    tables.append(
                        ContractTable(
                            headers=headers,
                            rows=rows,
                            slide_index=slide_index,
                            source="pptx",
                        )
                    )
        slide_texts[slide_index] = "\n".join(current_slide_texts)
    return "\n".join(texts), tables, slide_texts


def _unpack_pptx_extraction(
    extraction: Any,
) -> tuple[str, list[ContractTable], dict[int, str]]:
    if len(extraction) == 2:
        text, tables = extraction
        return text, tables, {}
    return extraction


def _table_values_are_visible_text(table: ContractTable, text: str) -> bool:
    normalized_lines = Counter(
        normalized
        for line in text.splitlines()
        if (normalized := _norm(line))
    )
    values = [
        *table.headers,
        *(cell for row in table.rows for cell in row),
    ]
    for value in values:
        normalized = _norm(value)
        if not normalized:
            continue
        if normalized_lines[normalized] < 1:
            return False
        normalized_lines[normalized] -= 1
    return True


def _table_matches_expected_slide(
    expected: ContractTable,
    actual: ContractTable,
) -> bool:
    return expected.slide_index is None or actual.slide_index == expected.slide_index


def _pptx_text_for_table(
    expected: ContractTable,
    full_text: str,
    slide_texts: dict[int, str],
) -> str:
    if expected.slide_index is None:
        return full_text
    if expected.slide_index in slide_texts:
        return slide_texts[expected.slide_index]
    if slide_texts:
        return ""
    return full_text


def validate_pptx_contract(
    state: GenerationContractState,
    path: str,
) -> list[ContractIssue]:
    if not state.enabled or not path:
        return []
    text, tables, slide_texts = _unpack_pptx_extraction(_extract_pptx(path))
    issues: list[ContractIssue] = _generic_placeholder_issues(
        text.splitlines(),
        stage="pptx_export",
    )
    for locked in state.locked_text:
        if locked not in text:
            issues.append(
                ContractIssue(
                    reason="missing_locked_text",
                    message="Exported PPTX did not preserve locked text verbatim.",
                    stage="pptx_export",
                    expected=locked,
                )
            )
    for term in state.exact_terms:
        if term and term not in text:
            issues.append(
                ContractIssue(
                    reason="changed_metric_date_or_label",
                    message="Exported PPTX is missing an exact metric, date, or label.",
                    stage="pptx_export",
                    expected=term,
                )
            )
    normalized_text = _norm(text)
    for forbidden in state.forbidden_additions:
        if _norm(forbidden) and _norm(forbidden) in normalized_text:
            issues.append(
                ContractIssue(
                    reason="forbidden_addition",
                    message="Exported PPTX contains forbidden content.",
                    stage="pptx_export",
                    expected=forbidden,
                )
            )
    for expected in state.evidence_tables:
        table_matches = any(
            _table_matches(expected, actual)
            and _table_matches_expected_slide(expected, actual)
            for actual in tables
        )
        table_values_visible = _table_values_are_visible_text(
            expected,
            _pptx_text_for_table(expected, text, slide_texts),
        )
        if table_matches:
            continue
        # The current PPTX exporter can flatten HTML table layouts into visible
        # positioned text/shapes. Strict export validation cares that evidence
        # values survived visibly; native table objects are a stronger pass.
        if not table_values_visible:
            issues.append(
                ContractIssue(
                    reason="changed_table_values",
                    message="Exported PPTX did not preserve expected evidence table values.",
                    stage="pptx_export",
                    section_index=expected.slide_index,
                    expected={"headers": expected.headers, "rows": expected.rows},
                )
            )
    return issues
